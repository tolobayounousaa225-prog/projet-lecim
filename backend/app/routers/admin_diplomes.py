"""Impression des diplômes pour le chargé d'examen : import Excel des candidats
admis, génération en lot (un PDF par ligne, à la manière d'un publipostage),
téléchargement groupé en ZIP. Même logique de confiance que les cartes et
attestations (numéro imprévisible, QR de vérification publique)."""

import datetime
import io
import secrets
import zipfile
from pathlib import Path

from fastapi import APIRouter, Depends, Form, Request, UploadFile, status
from fastapi.responses import RedirectResponse, Response
from fastapi.templating import Jinja2Templates
from sqlalchemy import func
from sqlalchemy.orm import Session

from .. import audit, models
from ..attestations_pdf import generate_diplome_pdf
from ..database import get_db
from ..deps import require_diplomes_access_web

router = APIRouter(prefix="/admin/diplomes", tags=["admin-diplomes"])

templates_dir = Path(__file__).resolve().parent.parent / "templates"
templates = Jinja2Templates(directory=str(templates_dir))

# Colonnes attendues dans le fichier Excel importé (en-têtes reconnus, insensibles
# à la casse et aux accents) — doivent correspondre au modèle téléchargeable.
COLUMN_ALIASES = {
    "nom complet": "nom_complet",
    "nom et prenoms": "nom_complet",
    "nom et prenom": "nom_complet",
    "etablissement": "etablissement_nom",
    "type d'examen": "type_examen",
    "type examen": "type_examen",
    "examen": "type_examen",
    "mention": "mention",
    "resultat": "mention",
    "mention / resultat": "mention",
    "mention/resultat": "mention",
    "annee scolaire": "annee_scolaire",
    "date de naissance": "date_naissance",
    "ne(e) le": "date_naissance",
    "date naissance": "date_naissance",
}
REQUIRED_FIELDS = ["nom_complet", "etablissement_nom", "type_examen", "annee_scolaire", "date_naissance"]
FIELD_LABELS = {
    "nom_complet": "Nom complet",
    "etablissement_nom": "Établissement",
    "type_examen": "Type d'examen",
    "annee_scolaire": "Année scolaire",
    "date_naissance": "Date de naissance",
}


def _normalize_header(value) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    for a, b in (("é", "e"), ("è", "e"), ("ê", "e"), ("à", "a"), ("ô", "o"), ("î", "i"), ("ç", "c")):
        text = text.replace(a, b)
    return text


def _parse_date_cell(value) -> datetime.date | None:
    """La cellule peut arriver déjà typée (Excel formate la colonne en date) ou
    sous forme de texte (DD/MM/YYYY le plus courant en saisie manuelle, ou
    YYYY-MM-DD) — on tente les deux avant d'abandonner."""
    if value is None:
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _detect_duplicates(db: Session, candidates: list[dict]) -> list[str]:
    """Repère les candidats qui semblent déjà avoir un diplôme pour le même
    examen — même nom complet (insensible à la casse/aux espaces) et même date
    de naissance, pour le même type d'examen. Compare les lignes entre elles
    (doublon dans le fichier lui-même) ET contre les diplômes déjà en base
    (ré-import accidentel d'un lot déjà traité). Ne bloque jamais seul :
    affiché comme avertissement, à confirmer explicitement pour continuer."""
    seen: set[tuple[str, object, str]] = set()
    warnings: list[str] = []

    for values in candidates:
        key = (values["nom_complet"].strip().lower(), values["date_naissance"], values["type_examen"].strip().lower())
        if key in seen:
            warnings.append(
                f"« {values['nom_complet']} » ({values['type_examen']}) apparaît en double dans le fichier importé."
            )
        else:
            seen.add(key)

    noms = {v["nom_complet"].strip().lower() for v in candidates}
    if noms:
        existants = (
            db.query(models.Diplome.nom_complet, models.Diplome.date_naissance, models.Diplome.type_examen)
            .filter(func.lower(func.trim(models.Diplome.nom_complet)).in_(noms))
            .all()
        )
        existants_set = {(e.nom_complet.strip().lower(), e.date_naissance, e.type_examen.strip().lower()) for e in existants}
        for values in candidates:
            key = (values["nom_complet"].strip().lower(), values["date_naissance"], values["type_examen"].strip().lower())
            if key in existants_set:
                warnings.append(
                    f"« {values['nom_complet']} » a déjà un diplôme enregistré pour l'examen "
                    f"« {values['type_examen']} » avec la même date de naissance."
                )

    return warnings


def _next_numero_diplome(db: Session) -> str:
    """Numéro imprévisible (même logique que les cartes, voir admin_cartes.py) :
    un numéro devinable permettrait d'énumérer /verify-diplome/{numero} et
    d'accéder aux informations d'autres candidats."""
    for _ in range(10):
        candidate = f"LECIM-DIP-{secrets.token_hex(4).upper()}"
        exists = db.query(models.Diplome.id).filter(models.Diplome.numero_diplome == candidate).first()
        if not exists:
            return candidate
    raise RuntimeError("Impossible de générer un numéro de diplôme unique")


def _safe_filename(name: str, fallback: str) -> str:
    cleaned = "".join(c if c.isalnum() or c in " -_" else "_" for c in name).strip()
    return cleaned or fallback


def _build_zip(diplomes: list["models.Diplome"]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for diplome in diplomes:
            pdf_bytes = generate_diplome_pdf(diplome)
            filename = f"{_safe_filename(diplome.nom_complet, diplome.numero_diplome)} - {diplome.numero_diplome}.pdf"
            zf.writestr(filename, pdf_bytes)
    return buffer.getvalue()


@router.get("")
def diplomes_list(
    request: Request,
    db: Session = Depends(get_db),
    user: models.User = Depends(require_diplomes_access_web),
):
    items = db.query(models.Diplome).order_by(models.Diplome.created_at.desc()).limit(1000).all()
    batches: dict[str, list] = {}
    for item in items:
        batches.setdefault(item.import_batch, []).append(item)
    return templates.TemplateResponse(
        request,
        "admin/diplomes_list.html",
        {"admin": user, "batches": batches, "active": "diplomes"},
    )


@router.get("/statistiques")
def diplomes_statistiques(
    request: Request,
    db: Session = Depends(get_db),
    user: models.User = Depends(require_diplomes_access_web),
):
    totaux = (
        db.query(
            models.Diplome.annee_scolaire,
            models.Diplome.type_examen,
            func.count(models.Diplome.id).label("total"),
        )
        .group_by(models.Diplome.annee_scolaire, models.Diplome.type_examen)
        .order_by(models.Diplome.annee_scolaire.desc(), models.Diplome.type_examen)
        .all()
    )
    mentions_rows = (
        db.query(
            models.Diplome.annee_scolaire,
            models.Diplome.type_examen,
            models.Diplome.mention,
            func.count(models.Diplome.id).label("total"),
        )
        .filter(models.Diplome.mention.isnot(None))
        .group_by(models.Diplome.annee_scolaire, models.Diplome.type_examen, models.Diplome.mention)
        .all()
    )
    mentions_by_group: dict[tuple[str, str], list[tuple[str, int]]] = {}
    for row in mentions_rows:
        key = (row.annee_scolaire, row.type_examen)
        mentions_by_group.setdefault(key, []).append((row.mention, row.total))

    groupes = [
        {
            "annee_scolaire": row.annee_scolaire,
            "type_examen": row.type_examen,
            "total": row.total,
            "mentions": sorted(mentions_by_group.get((row.annee_scolaire, row.type_examen), []), key=lambda m: -m[1]),
        }
        for row in totaux
    ]
    total_general = sum(g["total"] for g in groupes)

    return templates.TemplateResponse(
        request,
        "admin/diplomes_statistiques.html",
        {"admin": user, "groupes": groupes, "total_general": total_general, "active": "diplomes"},
    )


@router.get("/modele.xlsx")
def diplomes_modele_xlsx(
    user: models.User = Depends(require_diplomes_access_web),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    sheet = wb.active
    sheet.title = "Admis"

    headers = ["Nom complet", "Établissement", "Type d'examen", "Mention / Résultat", "Année scolaire", "Date de naissance"]
    sheet.append(headers)
    header_fill = PatternFill(start_color="0B3D2E", end_color="0B3D2E", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col in range(1, len(headers) + 1):
        sheet.cell(row=1, column=col).fill = header_fill
        sheet.cell(row=1, column=col).font = header_font

    sheet.append(["Fatou KONE", "Madrassa Al Houda", "CEPE", "Admis(e)", "2025-2026", "12/06/2015"])
    example_font = Font(italic=True, color="888888")
    for cell in sheet[2]:
        cell.font = example_font
    sheet.cell(row=2, column=6).number_format = "DD/MM/YYYY"

    widths = [28, 30, 20, 20, 16, 18]
    for i, w in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + i)].width = w

    buffer = io.BytesIO()
    wb.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="modele_diplomes_lecim.xlsx"'},
    )


@router.get("/importer")
def diplomes_importer_form(
    request: Request,
    user: models.User = Depends(require_diplomes_access_web),
):
    return templates.TemplateResponse(
        request,
        "admin/diplomes_importer.html",
        {
            "admin": user, "active": "diplomes", "errors": None, "warnings": None,
            "today": datetime.date.today().isoformat(),
        },
    )


@router.post("/importer")
async def diplomes_importer_process(
    request: Request,
    fichier: UploadFile,
    date_delivrance: str = Form(...),
    confirmer_doublons: bool = Form(False),
    db: Session = Depends(get_db),
    user: models.User = Depends(require_diplomes_access_web),
):
    from openpyxl import load_workbook

    def _error(messages: list[str], warnings: list[str] | None = None):
        return templates.TemplateResponse(
            request,
            "admin/diplomes_importer.html",
            {
                "admin": user, "active": "diplomes", "errors": messages, "warnings": warnings,
                "today": date_delivrance,
            },
            status_code=status.HTTP_400_BAD_REQUEST,
        )

    try:
        date_livr = datetime.datetime.strptime(date_delivrance, "%Y-%m-%d").date()
    except ValueError:
        return _error(["Date de délivrance invalide."])

    raw = await fichier.read()
    if not raw:
        return _error(["Le fichier envoyé est vide."])

    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = wb.active
    except Exception:
        return _error(["Fichier illisible : vérifiez qu'il s'agit bien d'un fichier Excel (.xlsx)."])

    rows_iter = sheet.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        return _error(["Le fichier ne contient aucune ligne d'en-tête."])

    column_index: dict[str, int] = {}
    for idx, raw_header in enumerate(header_row):
        key = COLUMN_ALIASES.get(_normalize_header(raw_header))
        if key and key not in column_index:
            column_index[key] = idx

    missing = [f for f in REQUIRED_FIELDS if f not in column_index]
    if missing:
        return _error([
            "Colonnes manquantes dans le fichier : " + ", ".join(FIELD_LABELS[f] for f in missing) + ". "
            "Utilisez le modèle Excel téléchargeable ci-dessus, sans renommer les en-têtes."
        ])

    candidates = []
    row_errors = []
    for row_number, row in enumerate(rows_iter, start=2):
        if row is None or all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        values = {}
        for field, idx in column_index.items():
            cell = row[idx] if idx < len(row) else None
            if field == "date_naissance":
                values[field] = _parse_date_cell(cell)
            else:
                values[field] = str(cell).strip() if cell is not None else ""
        row_missing = [FIELD_LABELS[f] for f in REQUIRED_FIELDS if not values.get(f)]
        if row_missing:
            row_errors.append(f"Ligne {row_number} : champ(s) manquant(s) ou illisible(s) — {', '.join(row_missing)}.")
            continue
        candidates.append(values)

    if row_errors:
        return _error(row_errors)
    if not candidates:
        return _error(["Aucune ligne de candidat trouvée dans le fichier."])

    if not confirmer_doublons:
        duplicate_warnings = _detect_duplicates(db, candidates)
        if duplicate_warnings:
            return _error([], warnings=duplicate_warnings)

    import_batch = f"LOT-{datetime.datetime.utcnow().strftime('%Y%m%d%H%M%S')}-{secrets.token_hex(2)}"
    generated = []
    for values in candidates:
        diplome = models.Diplome(
            numero_diplome=_next_numero_diplome(db),
            nom_complet=values["nom_complet"],
            etablissement_nom=values["etablissement_nom"],
            type_examen=values["type_examen"],
            mention=values.get("mention") or None,
            annee_scolaire=values["annee_scolaire"],
            date_naissance=values["date_naissance"],
            date_delivrance=date_livr,
            import_batch=import_batch,
            generated_by_id=user.id,
        )
        db.add(diplome)
        db.flush()
        generated.append(diplome)

    audit.log(db, user, "create", "Diplômes", None, f"A généré {len(generated)} diplôme(s) (lot {import_batch})")
    db.commit()

    zip_bytes = _build_zip(generated)
    return Response(
        content=zip_bytes,
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="diplomes_{import_batch}.zip"'},
    )


@router.get("/{diplome_id}/pdf")
def diplome_pdf(
    diplome_id: int,
    db: Session = Depends(get_db),
    user: models.User = Depends(require_diplomes_access_web),
):
    diplome = db.get(models.Diplome, diplome_id)
    if not diplome:
        return RedirectResponse(url="/admin/diplomes", status_code=status.HTTP_303_SEE_OTHER)
    pdf_bytes = generate_diplome_pdf(diplome)
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="diplome_{diplome.numero_diplome}.pdf"'},
    )


@router.get("/lot/{import_batch}.zip")
def diplomes_batch_zip(
    import_batch: str,
    db: Session = Depends(get_db),
    user: models.User = Depends(require_diplomes_access_web),
):
    items = db.query(models.Diplome).filter(models.Diplome.import_batch == import_batch).all()
    if not items:
        return RedirectResponse(url="/admin/diplomes", status_code=status.HTTP_303_SEE_OTHER)
    zip_bytes = _build_zip(items)
    return Response(
        content=zip_bytes,
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="diplomes_{import_batch}.zip"'},
    )
